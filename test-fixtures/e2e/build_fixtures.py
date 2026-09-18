from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent


def create_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("IOM 021 2026 Perubahan Ketentuan Cuti Tahunan")

    opening = document.add_paragraph()
    opening.add_run("Kepada seluruh karyawan PT Contoh Nusantara").bold = True
    document.add_heading("Ketentuan yang berlaku", level=1)
    document.add_paragraph(
        "Mulai 1 Februari 2026, karyawan yang telah bekerja 12 bulan berturut-turut memperoleh "
        "12 hari kerja cuti tahunan "
        "dalam satu tahun kalender. Permohonan cuti diajukan melalui portal HR paling lambat tiga "
        "hari kerja sebelum tanggal cuti. Maksimal 5 hari dapat dibawa ke tahun berikutnya dan harus "
        "digunakan paling lambat 30 Juni."
    )
    document.add_heading("Hubungan dengan aturan sebelumnya", level=1)
    document.add_paragraph(
        "IOM ini diusulkan sebagai pengganti IOM 014 2014 untuk topik hak cuti tahunan. Hubungan "
        "penggantian baru berlaku setelah HR mengonfirmasinya di sistem. Dokumen uji ketiga ini "
        "memakai pemisahan section untuk menguji sanitasi parsial."
    )
    document.add_heading("Lampiran terbatas untuk HR", level=1)
    document.add_paragraph(
        "Simulasi internal memperkirakan tambahan biaya tenaga pengganti sebesar Rp842.750.000. "
        "Anggaran cadangan Divisi Operasional adalah Rp1.125.000.000 dan kode pusat biaya CC-HR-9917. "
        "Nilai, kode, dan simulasi ini hanya boleh diakses HR dan tidak boleh ditampilkan kepada karyawan."
    )

    styles = document.styles
    for style_name in ("Normal", "Title", "Heading 1"):
        style = styles[style_name]
        style.font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(20)
    styles["Heading 1"].font.size = Pt(13)
    document.save(ROOT / "iom-021-2026-cuti-baru.docx")


def create_scan_pdf() -> None:
    width, height = 1654, 2339
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
    ]
    font_path = next((path for path in font_candidates if path.exists()), None)
    title_font = ImageFont.truetype(str(font_path), 46) if font_path else ImageFont.load_default()
    body_font = ImageFont.truetype(str(font_path), 30) if font_path else ImageFont.load_default()
    y = 150
    draw.text((150, y), "IOM 029 2026 Prosedur Keadaan Darurat", fill="black", font=title_font)
    y += 110
    paragraphs = [
        "Mulai 20 September 2026, seluruh karyawan wajib mengikuti arahan petugas lantai saat alarm berbunyi.",
        "Gunakan tangga darurat, jangan menggunakan lift, dan berkumpul di titik kumpul timur gedung.",
        "Atasan melakukan pencatatan kehadiran setelah evakuasi. Karyawan yang membutuhkan bantuan mobilitas harus menghubungi HR sebelum simulasi berlangsung.",
        "Nomor darurat internal adalah 1190. Dokumen ini dapat dibaca seluruh karyawan.",
    ]
    for paragraph in paragraphs:
        for line in wrap(paragraph, width=72):
            draw.text((150, y), line, fill="black", font=body_font)
            y += 48
        y += 36
    image.save(ROOT / "iom-029-2026-darurat-scan.pdf", "PDF", resolution=150.0)


if __name__ == "__main__":
    import sys

    target = sys.argv[1]
    if target == "docx":
        create_docx()
    elif target == "pdf":
        create_scan_pdf()
    else:
        raise SystemExit("target must be docx or pdf")
